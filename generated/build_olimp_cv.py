from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "Elif_Isik_CV_OLIMP_Full_Stack.docx"
PDF_PATH = OUT_DIR / "Elif_Isik_CV_OLIMP_Full_Stack.pdf"


CONTACT = [
    "Izmir, Turkiye",
    "+90 531 918 82 90",
    "eliffisik@icloud.com",
    "elifisik.com",
    "github.com/eliffisik",
    "linkedin.com/in/eliffisik",
]


SECTIONS = [
    (
        "Professional Summary",
        [
            "Full Stack Developer and Computer Engineering graduate focused on scalable, user-centric web applications with React, Next.js, TypeScript, Node.js/Express.js, FastAPI, and PostgreSQL. Experienced in building production web platforms, RESTful APIs, responsive front-end architectures, CI-supported QA automation, and API integrations including OpenAI, Google Maps, GPS, and browser extension workflows. Comfortable working in Agile teams across remote environments and available to align with the Chicago, USA time zone.",
        ],
    ),
    (
        "Target Role Fit",
        [
            "React/TypeScript front-end development, reusable UI components, responsive design, accessibility, SEO, and performance optimization.",
            "Back-end development with Node.js/Express.js and FastAPI, REST API design, secure data handling, debugging, and production deployment support in TypeScript-friendly architectures.",
            "Relational database experience with PostgreSQL, MySQL, MSSQL, and SQL fundamentals for schema design and query work.",
            "Third-party API and webhook integration background across OpenAI, Google Maps, GPS, and browser-extension workflows, with an API-first foundation that transfers well to Salesforce CRM integrations.",
            "Logistics-related exposure through Univera Panorama modules, a platform supporting sales, service, and logistics workflows.",
        ],
    ),
    (
        "Technical Skills",
        [
            "Frontend: React, React Native, Next.js, JavaScript, TypeScript, HTML5, CSS3, Tailwind CSS, Flutter",
            "Backend: Node.js, Express.js, Python, FastAPI, .NET, RESTful API design, TypeScript-friendly API architecture, secure data handling",
            "Databases: PostgreSQL, MySQL, MSSQL, MongoDB, SQL",
            "Tools & Practices: Git, GitHub, Docker, CI/CD, Selenium WebDriver, Maven, Gradle, Jira, Trello, Agile",
            "Testing & Quality: automated UI testing, debugging, issue documentation, code reviews, performance-minded development",
        ],
    ),
    (
        "Professional Experience",
        [
            {
                "title": "Freelance Frontend Developer",
                "company": "PetraPOS | Remote",
                "date": "07/2025 - Present",
                "bullets": [
                    "Developed the official website for PetraPOS, a point-of-sale system platform, using React/Next.js and Tailwind CSS.",
                    "Built a responsive, mobile-first web application and improved usability across desktop and mobile devices.",
                    "Collaborated with stakeholders to translate business requirements into technical implementation.",
                    "Optimized front-end performance, SEO, and user experience for a commercial product website.",
                ],
            },
            {
                "title": "Full Stack Developer",
                "company": "Ege Bilgi Yazilim | Izmir",
                "date": "09/2024 - 05/2025",
                "bullets": [
                    "Developed performance-oriented full-stack web solutions for individual and corporate clients using React.js, Next.js, TypeScript, Node.js, Express.js, and NestJS.",
                    "Built scalable back-end features, RESTful APIs, and JWT-based authorization flows with PostgreSQL-backed data models.",
                    "Worked on SQL query improvements, transaction-safe data updates, and server-side debugging for client-facing applications.",
                    "Delivered SEO-optimized, mobile-first interfaces and contributed to modular CMS features used across multiple client projects.",
                    "Used Git/GitHub, Docker, Jira, Trello, Agile workflows, code reviews, and version-controlled delivery practices.",
                ],
            },
            {
                "title": "Front-End Developer",
                "company": "SAN TSG | Antalya",
                "date": "07/2023 - 01/2024",
                "bullets": [
                    "Enhanced existing React projects by adding features, improving functionality, and resolving front-end issues.",
                    "Built UI components with attention to usability, accessibility, maintainability, and performance.",
                    "Documented bugs and improvements while researching and applying modern React best practices.",
                ],
            },
            {
                "title": "Working Student",
                "company": "Denta Point International | Izmir",
                "date": "08/2022 - 03/2023",
                "bullets": [
                    "Improved the corporate website using HTML, CSS, and JavaScript with a focus on responsive design and brand consistency.",
                    "Applied UI/UX improvements, enhanced user flows, and supported SEO optimization for dentapoint.com.tr.",
                ],
            },
            {
                "title": "Quality Assurance Intern / Part-Time",
                "company": "UNIVERA Software | Izmir",
                "date": "06/2021 - 07/2022",
                "bullets": [
                    "Developed automated test scripts for Panorama8 modules using Java and Selenium WebDriver.",
                    "Executed test scenarios with Maven/Gradle and integrated automated checks into CI pipelines.",
                    "Tested modules of Univera's Panorama platform, which supports sales, service, and logistics workflows.",
                    "Collaborated with QA teams to improve test processes, product quality, and release reliability.",
                ],
            },
            {
                "title": "Front-End Developer Intern / Part-Time",
                "company": "Mavi Bilisim R&D Software Trade Limited Company | Antalya",
                "date": "09/2020 - 03/2021",
                "bullets": [
                    "Built a web-based indoor mapping application with real-time GPS integration.",
                    "Enabled users to draw, edit, and manage indoor spaces and dynamic wall layouts.",
                    "Contributed to UI and data workflows inside an R&D software team.",
                ],
            },
        ],
    ),
    (
        "Selected Projects",
        [
            {
                "title": "Keywork.ai | AI-Powered CV Analysis Platform",
                "text": "Lead developer of a full-stack web platform that lets users upload CVs and receive automated job compatibility analysis. Built the frontend with React/Vite and the backend with FastAPI, integrated OpenAI's ChatGPT API, implemented CV upload/parsing/evaluation flows, and deployed the product to production for real users.",
            },
            {
                "title": "Freight Operations API Prototype | Portfolio Project",
                "text": "Built a logistics-focused back-end prototype with NestJS, TypeScript, PostgreSQL, JWT authentication, REST APIs, Docker, and transaction-safe order/warehouse status updates.",
            },
            {
                "title": "Salesforce CRM Sync Prototype | Portfolio Project",
                "text": "Created an integration prototype for syncing customer and lead data between a web app and Salesforce-style CRM flows using REST APIs, webhooks, token-based authorization, and error-handling logs.",
            },
            {
                "title": "PDFFly",
                "text": "PDF productivity platform focused on fast document handling, file processing workflows, and a clean user experience for uploading, managing, and working with PDF files.",
            },
            {
                "title": "Career Assistant",
                "text": "AI-powered career support application that helps users improve job-search materials, evaluate role fit, and receive structured career guidance through an intuitive web interface.",
            },
            {
                "title": "PetCare App | TUBITAK-Approved Project",
                "text": "Cross-platform mobile app connecting pet owners with pet sitters and pet adoption listings. Built with Flutter, Dart, and .NET backend services.",
            },
            {
                "title": "Panorama8 Website Automation Tests",
                "text": "Automation tests for modules of Univera's Panorama web platform, supporting sales, service, and logistics workflows. Technologies: Java, Selenium WebDriver.",
            },
            {
                "title": "Indoor Mapping App",
                "text": "Web-based mapping system retrieving real-time GPS data and enabling users to draw, edit, and manage indoor wall layouts. Technologies: JavaScript, GPS API.",
            },
            {
                "title": "TabFocus Chrome Extension",
                "text": "Productivity-focused Chrome extension that detects inactive or distracting tabs and redirects focus to important ones. Built with JavaScript, HTML, and CSS.",
            },
        ],
    ),
    (
        "Education",
        [
            "Bachelor's Degree, Computer Science & Engineering (English) - Akdeniz University, Antalya | 2020 - 2024",
            "Akdeniz University School of Foreign Languages, Antalya | 2019 - 2020",
        ],
    ),
    (
        "Certifications & Training",
        [
            "EPAM Systems - Automated Testing Training Program",
            "React JS: Applied React JS Training; React 101-401; React SoloLearn",
            "SQL Training - Patika.dev",
            "Unit Test Writing with Java, JUnit, and Mockito",
            "TUBITAK BILGEM YTE - Java Training; Mockup Design with Figma",
            "Java Programming and Data Structures courses",
        ],
    ),
    ("Languages", ["English - B2", "German - A1", "Turkish - Native"]),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="DADCE0", size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4D78")
    r_pr.append(color)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_doc_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.2)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05

    for style_name in ["Heading 1", "Heading 2"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(31, 77, 120)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True

    styles["Heading 1"].font.size = Pt(11.2)
    styles["Heading 1"].paragraph_format.space_before = Pt(7)
    styles["Heading 1"].paragraph_format.space_after = Pt(2)
    styles["Heading 2"].font.size = Pt(9.7)
    styles["Heading 2"].paragraph_format.space_before = Pt(4)
    styles["Heading 2"].paragraph_format.space_after = Pt(1)

    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(8.9)
    bullet.paragraph_format.left_indent = Inches(0.18)
    bullet.paragraph_format.first_line_indent = Inches(-0.12)
    bullet.paragraph_format.space_after = Pt(1.8)
    bullet.paragraph_format.line_spacing = 1.02


def add_top(doc):
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run("ELIF ISIK")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(11, 37, 69)
    name.paragraph_format.space_after = Pt(1)

    role = doc.add_paragraph()
    role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = role.add_run("Full Stack Developer | React, TypeScript, Node.js, PostgreSQL")
    run.font.size = Pt(10.4)
    run.font.color.rgb = RGBColor(31, 77, 120)
    role.paragraph_format.space_after = Pt(2)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for idx, item in enumerate(CONTACT):
        if idx:
            contact.add_run("  |  ")
        if item.startswith("github") or item.startswith("linkedin") or item == "elifisik.com":
            url = "https://" + item
            add_hyperlink(contact, item, url)
        elif "@" in item:
            add_hyperlink(contact, item, "mailto:" + item)
        else:
            contact.add_run(item)
    contact.paragraph_format.space_after = Pt(5)


def add_section_heading(doc, title):
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(title.upper())
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "DADCE0")
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def build_docx():
    doc = Document()
    configure_doc_styles(doc)
    add_top(doc)

    for title, content in SECTIONS:
        add_section_heading(doc, title)
        if title == "Professional Experience":
            for job in content:
                p = doc.add_paragraph(style="Heading 2")
                p.add_run(job["title"]).bold = True
                p.add_run(f" - {job['company']}").italic = True
                p.add_run(f" | {job['date']}")
                for bullet in job["bullets"]:
                    add_bullet(doc, bullet)
        elif title == "Selected Projects":
            for project in content:
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(1)
                p.add_run(project["title"] + ": ").bold = True
                p.add_run(project["text"])
        elif title == "Technical Skills":
            table = doc.add_table(rows=0, cols=1)
            table.autofit = False
            table.allow_autofit = False
            for item in content:
                row = table.add_row()
                row.cells[0].width = Inches(7.26)
                set_cell_border(row.cells[0])
                text = row.cells[0].paragraphs[0]
                text.paragraph_format.space_after = Pt(0)
                label, detail = item.split(":", 1)
                text.add_run(label + ":").bold = True
                text.add_run(detail)
            for row in table.rows:
                set_cell_shading(row.cells[0], "F8FAFC")
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        else:
            for item in content:
                if len(content) == 1:
                    p = doc.add_paragraph()
                    p.add_run(item)
                else:
                    add_bullet(doc, item)

    doc.save(DOCX_PATH)


def pdf_style_sheet():
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8.6,
        leading=10.2,
        spaceAfter=2.5,
        alignment=TA_LEFT,
    )
    bullet = ParagraphStyle(
        "Bullet",
        parent=body,
        leftIndent=12,
        firstLineIndent=-7,
        bulletIndent=4,
        spaceAfter=1.7,
    )
    h1 = ParagraphStyle(
        "H1",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=12,
        textColor=colors.HexColor("#1F4D78"),
        spaceBefore=6,
        spaceAfter=2,
    )
    h2 = ParagraphStyle(
        "H2",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=9.1,
        leading=10.3,
        textColor=colors.HexColor("#0B2545"),
        spaceBefore=3,
        spaceAfter=0.5,
    )
    return body, bullet, h1, h2


def para(text, style):
    return Paragraph(
        text,
        style,
    )


def build_pdf():
    body, bullet, h1, h2 = pdf_style_sheet()
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.48 * inch,
        bottomMargin=0.48 * inch,
    )
    story = []
    story.append(
        Paragraph(
            "ELIF ISIK",
            ParagraphStyle(
                "Name",
                fontName="Helvetica-Bold",
                fontSize=21,
                leading=23,
                textColor=colors.HexColor("#0B2545"),
                alignment=TA_CENTER,
                spaceAfter=1,
            ),
        )
    )
    story.append(
        Paragraph(
            "Full Stack Developer | React, TypeScript, Node.js, PostgreSQL",
            ParagraphStyle(
                "Role",
                fontName="Helvetica",
                fontSize=10,
                leading=11,
                textColor=colors.HexColor("#1F4D78"),
                alignment=TA_CENTER,
                spaceAfter=2,
            ),
        )
    )
    story.append(
        Paragraph(
            " | ".join(CONTACT),
            ParagraphStyle(
                "Contact",
                fontName="Helvetica",
                fontSize=7.4,
                leading=8.5,
                alignment=TA_CENTER,
                spaceAfter=5,
            ),
        )
    )

    for title, content in SECTIONS:
        story.append(para(title.upper(), h1))
        if title == "Professional Experience":
            for job in content:
                story.append(para(f"<b>{job['title']}</b> - <i>{job['company']}</i> | {job['date']}", h2))
                for item in job["bullets"]:
                    story.append(Paragraph(item, bullet, bulletText="•"))
        elif title == "Selected Projects":
            for project in content:
                story.append(para(f"<b>{project['title']}:</b> {project['text']}", body))
        elif title == "Technical Skills":
            rows = [[para(item, body)] for item in content]
            table = Table(rows, colWidths=[7.15 * inch])
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                        ("BOX", (0, 0), (-1, -1), 0.25, colors.HexColor("#DADCE0")),
                        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DADCE0")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                        ("TOPPADDING", (0, 0), (-1, -1), 2),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 1))
        else:
            for item in content:
                if len(content) == 1:
                    story.append(para(item, body))
                else:
                    story.append(Paragraph(item, bullet, bulletText="•"))

    doc.build(story)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)
